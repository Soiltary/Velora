"""Local conversion engine. No network operations; originals are never overwritten."""
from __future__ import annotations

import csv
import html
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from PIL import Image, ImageOps, ImageSequence

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except ImportError:
    pass

IMAGES = {'.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff', '.gif', '.ico', '.heic', '.heif', '.avif'}
AUDIO = {'.mp3', '.wav', '.flac', '.m4a', '.aac', '.ogg', '.opus', '.aiff', '.wma'}
VIDEO = {'.mp4', '.mkv', '.mov', '.webm', '.avi', '.wmv', '.m4v'}
OFFICE = {'.doc', '.docx', '.odt', '.rtf', '.ppt', '.pptx', '.odp', '.xls', '.xlsx', '.ods'}


def actions(path: Path, tools=False) -> list[str]:
    ext = path.suffix.lower()
    if path.is_dir():
        return ['zip']
    if tools:
        if ext in IMAGES:
            return ['compress', 'resize', 'rotate', 'grayscale', 'strip-metadata']
        if ext == '.pdf':
            return ['compress', 'split-pages']
        if ext in VIDEO:
            return ['compress', 'trim', 'extract-audio']
        if ext in AUDIO:
            return ['trim']
        return ['zip']
    if ext in IMAGES:
        return ['png', 'jpg', 'webp', 'avif', 'tiff', 'bmp', 'ico', 'pdf']
    if ext in VIDEO:
        return ['mp4', 'mkv', 'webm', 'mov', 'gif', 'mp3', 'wav']
    if ext in AUDIO:
        return ['mp3', 'wav', 'flac', 'm4a', 'ogg', 'opus', 'aiff']
    if ext == '.pdf':
        return ['png', 'jpg', 'docx', 'txt']
    if ext in {'.srt', '.vtt'}:
        return ['srt', 'vtt', 'txt']
    if ext in {'.csv', '.tsv', '.xlsx'}:
        return ['csv', 'tsv', 'xlsx', 'pdf']
    if ext in {'.txt', '.md', '.html', '.htm', '.docx'}:
        return ['txt', 'html', 'docx', 'pdf']
    if ext in OFFICE:
        return ['pdf']
    return ['zip']


def common_actions(paths, tools=False):
    if not paths:
        return []
    first = actions(Path(paths[0]), tools)
    return [a for a in first if all(a in actions(Path(p), tools) for p in paths[1:])]


def libreoffice():
    candidates = [shutil.which('soffice'), shutil.which('libreoffice')]
    if os.name == 'nt':
        candidates += [str(Path(os.environ.get(k, 'C:/Program Files')) / 'LibreOffice/program/soffice.exe')
                       for k in ('ProgramFiles', 'ProgramFiles(x86)')]
    return next((p for p in candidates if p and Path(p).is_file()), None)


def ffmpeg():
    bundled = Path(__file__).resolve().parent.parent / 'vendor' / 'ffmpeg.exe'
    if bundled.is_file():
        return str(bundled)
    if os.name != 'nt' and shutil.which('ffmpeg'):
        return shutil.which('ffmpeg')
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


def run(args, cancel=None, timeout=1800):
    import time
    with tempfile.TemporaryFile() as log:
        proc = subprocess.Popen([str(a) for a in args], stdout=log, stderr=log,
                                creationflags=0x08000000 if os.name == 'nt' else 0)
        started = time.monotonic()
        while proc.poll() is None:
            if (cancel and cancel.is_set()) or time.monotonic() - started > timeout:
                proc.kill()
                proc.wait()
                raise RuntimeError('Cancelled' if cancel and cancel.is_set() else 'Conversion timed out')
            time.sleep(.08)
        if proc.returncode:
            log.seek(0)
            detail = log.read().decode('utf-8', errors='replace')[-1800:]
            raise RuntimeError(detail or 'Converter failed')


def office_convert(src, dest, work, cancel):
    exe = libreoffice()
    if not exe:
        raise RuntimeError('This conversion needs LibreOffice. Install it from libreoffice.org, then retry.')
    profile = (work / 'office-profile').as_uri()
    run([exe, '-env:UserInstallation=' + profile, '--headless', '--convert-to', dest,
         '--outdir', work, src], cancel)
    result = work / (src.stem + '.' + dest)
    if not result.is_file():
        raise RuntimeError('LibreOffice did not produce an output. The input may be unsupported or protected.')
    return [result]


def _raster(src, target, out, opts):
    with Image.open(src) as original:
        im = ImageOps.exif_transpose(original)
        # Multi-frame inputs intentionally use the first frame for still-image exports.
        im.load()
        if target in {'compress', 'resize', 'rotate', 'grayscale', 'strip-metadata'}:
            ext = 'jpg' if src.suffix.lower() in {'.jpg', '.jpeg'} else 'png'
            name = out / (src.stem + '-' + target + '.' + ext)
            if target == 'resize':
                im.thumbnail((int(opts.get('width', 1920)), int(opts.get('height', 1080))), Image.Resampling.LANCZOS)
            if target == 'rotate':
                im = im.rotate(-90, expand=True)
            if target == 'grayscale':
                im = ImageOps.grayscale(im)
        else:
            ext = target
            name = out / (src.stem + '.' + ext)
        if ext in {'jpg', 'pdf', 'bmp'}:
            if im.mode in {'RGBA', 'LA'} or 'transparency' in im.info:
                rgba = im.convert('RGBA')
                base = Image.new('RGB', im.size, 'white')
                base.paste(rgba, mask=rgba.getchannel('A'))
                im = base
            else:
                im = im.convert('RGB')
        clean = Image.new(im.mode, im.size)
        clean.paste(im)
        options = {'quality': int(opts.get('quality', 85)), 'optimize': True} if ext in {'jpg', 'webp', 'avif'} else {}
        if ext == 'png':
            options = {'optimize': True, 'compress_level': 9}
        if ext == 'ico':
            clean.thumbnail((256, 256), Image.Resampling.LANCZOS)
            options = {'sizes': [(16, 16), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)]}
        clean.save(name, **options)
        return [name]


def _media(src, target, out, opts, cancel):
    ext = src.suffix.lower()
    dest_ext = ('mp4' if ext in VIDEO else 'm4a') if target in {'compress', 'trim'} else target
    if target == 'extract-audio':
        dest_ext = 'mp3'
    name = out / (src.stem + ('-' + target if target in {'compress', 'trim', 'extract-audio'} else '') + '.' + dest_ext)
    args = [ffmpeg(), '-nostdin', '-hide_banner', '-loglevel', 'error', '-y', '-i', src]
    if target == 'trim':
        start, end = float(opts.get('start', 0)), float(opts.get('end', 30))
        if end <= start or start < 0:
            raise ValueError('End must be greater than start; times are seconds.')
        args += ['-ss', str(start), '-t', str(end - start)]
    if dest_ext in {'mp3', 'wav', 'flac', 'm4a', 'ogg', 'opus', 'aiff'}:
        codec = {'mp3': 'libmp3lame', 'wav': 'pcm_s16le', 'flac': 'flac', 'm4a': 'aac',
                 'ogg': 'libvorbis', 'opus': 'libopus', 'aiff': 'pcm_s16be'}[dest_ext]
        args += ['-vn', '-c:a', codec]
        if dest_ext in {'mp3', 'm4a', 'ogg'}:
            args += ['-b:a', '192k']
    elif dest_ext == 'gif':
        args += ['-filter_complex', '[0:v]fps=12,scale=640:-1:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse', '-an']
    else:
        args += ['-map', '0:v:0', '-map', '0:a:0?']
        if dest_ext == 'webm':
            args += ['-c:v', 'libvpx-vp9', '-crf', '34', '-b:v', '0', '-c:a', 'libopus']
        else:
            args += ['-c:v', 'libx264', '-preset', 'medium', '-crf', '28' if target == 'compress' else '20',
                     '-vf', 'scale=ceil(iw/2)*2:ceil(ih/2)*2', '-pix_fmt', 'yuv420p', '-c:a', 'aac']
            if dest_ext in {'mp4', 'mov'}:
                args += ['-movflags', '+faststart']
    run(args + [name], cancel)
    return [name]


def _pdf(src, target, out, opts, cancel):
    import fitz
    if target == 'docx':
        from .documents import pdf_to_word
        return pdf_to_word(src, out, cancel)
    results = []
    with fitz.open(src) as doc:
        if doc.needs_pass:
            raise ValueError('Password-protected PDFs must be unlocked before conversion.')
        if target == 'txt':
            dest = out / (src.stem + '.txt')
            dest.write_text('\n\n'.join(p.get_text() for p in doc), encoding='utf-8')
            return [dest]
        if target == 'compress':
            dest = out / (src.stem + '-compressed.pdf')
            doc.save(dest, garbage=4, deflate=True, clean=True)
            return [dest]
        for i, page in enumerate(doc):
            if cancel and cancel.is_set():
                raise RuntimeError('Cancelled')
            dest = out / f'{src.stem}-page-{i+1:03d}.{("pdf" if target == "split-pages" else target)}'
            if target == 'split-pages':
                with fitz.open() as single:
                    single.insert_pdf(doc, from_page=i, to_page=i)
                    single.save(dest)
            else:
                pix = page.get_pixmap(dpi=int(opts.get('dpi', 150)), alpha=False)
                pix.save(dest)
            results.append(dest)
    return results


def _subtitle(src, target, out):
    text = src.read_text(encoding='utf-8-sig')
    lines = text.replace('\r\n', '\n').split('\n')
    blocks = re.split(r'\n\s*\n', '\n'.join(lines).strip())
    cues = []
    pattern = re.compile(r'((?:\d{2}:)?\d{2}:\d{2})[.,](\d{3})\s*-->\s*((?:\d{2}:)?\d{2}:\d{2})[.,](\d{3})')
    for block in blocks:
        rows = block.splitlines()
        for index, row in enumerate(rows):
            match = pattern.search(row)
            if match:
                a, b, c, d = match.groups()
                if a.count(':') == 1:
                    a = '00:' + a
                if c.count(':') == 1:
                    c = '00:' + c
                cues.append((a, b, c, d, '\n'.join(rows[index+1:])))
                break
    if not cues:
        raise ValueError('No valid subtitle cues found. Use a UTF-8 SRT or VTT file.')
    dest = out / (src.stem + '.' + target)
    if target == 'txt':
        data = '\n\n'.join(re.sub('<[^>]+>', '', c[4]) for c in cues)
    else:
        sep = ',' if target == 'srt' else '.'
        data = 'WEBVTT\n\n' if target == 'vtt' else ''
        for i, (a, b, c, d, body) in enumerate(cues, 1):
            data += (f'{i}\n' if target == 'srt' else '') + f'{a}{sep}{b} --> {c}{sep}{d}\n{body}\n\n'
    dest.write_text(data, encoding='utf-8')
    return [dest]


def _table(src, target, out, cancel):
    from openpyxl import Workbook, load_workbook
    if target == 'pdf':
        return office_convert(src, 'pdf', out, cancel)
    if src.suffix.lower() == '.xlsx':
        wb = load_workbook(src, read_only=True, data_only=True)
        try:
            # Every worksheet is exported, keeping separate files for delimited output.
            if target == 'xlsx':
                dest = out / (src.stem + '.xlsx')
                shutil.copy2(src, dest)
                return [dest]
            results = []
            for i, sheet in enumerate(wb.worksheets):
                suffix = f'-sheet-{i+1}' if len(wb.worksheets) > 1 else ''
                dest = out / (src.stem + suffix + '.' + target)
                with dest.open('w', newline='', encoding='utf-8-sig') as f:
                    csv.writer(f, delimiter='\t' if target == 'tsv' else ',').writerows(sheet.values)
                results.append(dest)
            return results
        finally:
            wb.close()
    with src.open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f, delimiter='\t' if src.suffix.lower() == '.tsv' else ','))
    dest = out / (src.stem + '.' + target)
    if target == 'xlsx':
        wb = Workbook()
        for row in rows:
            # Preserve text literally; CSV formula injection must not become an XLSX formula.
            wb.active.append(row)
            for cell in wb.active[wb.active.max_row]:
                cell.data_type = 's'
        wb.save(dest)
        wb.close()
    else:
        with dest.open('w', newline='', encoding='utf-8-sig') as f:
            csv.writer(f, delimiter='\t' if target == 'tsv' else ',').writerows(rows)
    return [dest]


def _text(src, target, out, cancel):
    from docx import Document
    if target == 'pdf':
        if src.suffix.lower() == '.docx' and not libreoffice():
            from .documents import docx_to_pdf
            return docx_to_pdf(src, out, cancel)
        return office_convert(src, 'pdf', out, cancel)
    if src.suffix.lower() == '.docx':
        if target == 'docx':
            dest = out / (src.stem + '.docx')
            shutil.copy2(src, dest)
            return [dest]
        doc = Document(src)
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += ''.join('\n' + '\n'.join('\t'.join(c.text for c in r.cells) for r in t.rows) for t in doc.tables)
    else:
        text = src.read_text(encoding='utf-8-sig')
        if src.suffix.lower() in {'.html', '.htm'}:
            from html.parser import HTMLParser
            class PlainText(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.parts = []
                    self.skip = 0
                def handle_starttag(self, tag, attrs):
                    if tag in {'script', 'style'}:
                        self.skip += 1
                    if tag in {'p', 'br', 'div', 'tr', 'h1', 'h2', 'li'}:
                        self.parts.append('\n')
                def handle_endtag(self, tag):
                    if tag in {'script', 'style'}:
                        self.skip = max(0, self.skip - 1)
                def handle_data(self, data):
                    if not self.skip:
                        self.parts.append(data)
            parser = PlainText()
            parser.feed(text)
            text = ''.join(parser.parts)
    dest = out / (src.stem + '.' + target)
    if target == 'docx':
        doc = Document()
        for line in text.splitlines():
            doc.add_paragraph(line)
        doc.save(dest)
    elif target == 'html':
        dest.write_text('<!doctype html><meta charset="utf-8"><pre>' + html.escape(text) + '</pre>', encoding='utf-8')
    else:
        dest.write_text(text, encoding='utf-8')
    return [dest]


def _zip(src, out, cancel):
    dest = out / (src.stem + '.zip')
    with zipfile.ZipFile(dest, 'w', zipfile.ZIP_DEFLATED) as z:
        if src.is_dir():
            for root, dirs, files in os.walk(src, followlinks=False):
                dirs[:] = [d for d in dirs if not (Path(root) / d).is_symlink()]
                if not files and not dirs:
                    z.writestr(str(Path(root).relative_to(src.parent)) + '/', '')
                for name in files:
                    if cancel and cancel.is_set():
                        raise RuntimeError('Cancelled')
                    p = Path(root) / name
                    if not p.is_symlink() and not p.resolve().is_relative_to(out.resolve()):
                        z.write(p, p.relative_to(src.parent))
        else:
            z.write(src, src.name)
    return [dest]


def _publish(files, destination):
    published = []
    try:
        for f in files:
            for i in range(10000):
                name = f.stem + (f' ({i})' if i else '') + f.suffix
                path = destination / name
                try:
                    fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
                    os.close(fd)
                except FileExistsError:
                    continue
                try:
                    os.replace(f, path)
                except BaseException:
                    path.unlink(missing_ok=True)
                    raise
                published.append(path)
                break
            else:
                raise RuntimeError('Too many output files with the same name')
        return published
    except BaseException:
        for p in published:
            p.unlink(missing_ok=True)
        raise


def convert(path, target, destination=None, options=None, cancel=None):
    src = Path(path).resolve()
    if not src.exists():
        raise FileNotFoundError(str(src))
    if target != 'zip' and target not in actions(src) + actions(src, True):
        raise ValueError(f'{target} is not supported for {src.suffix or "this folder"}')
    dest = Path(destination).resolve() if destination else src.parent / 'Velora Output'
    if src.is_dir() and dest.is_relative_to(src):
        raise ValueError('Choose an output folder outside the folder being archived.')
    dest.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.velora-', dir=dest) as tmp:
        out = Path(tmp)
        ext = src.suffix.lower()
        opts = options or {}
        if target == 'zip':
            files = _zip(src, out, cancel)
        elif ext in IMAGES:
            files = _raster(src, target, out, opts)
        elif ext in AUDIO | VIDEO:
            files = _media(src, target, out, opts, cancel)
        elif ext == '.pdf':
            files = _pdf(src, target, out, opts, cancel)
        elif ext in {'.srt', '.vtt'}:
            files = _subtitle(src, target, out)
        elif ext in {'.csv', '.tsv', '.xlsx'}:
            files = _table(src, target, out, cancel)
        elif ext in {'.txt', '.md', '.html', '.htm', '.docx'}:
            files = _text(src, target, out, cancel)
        else:
            files = office_convert(src, target, out, cancel)
        if cancel and cancel.is_set():
            raise RuntimeError('Cancelled')
        if not files or any(not f.is_file() or not f.stat().st_size for f in files):
            # Empty extracted PDF text is a useful signal for scanned documents.
            raise RuntimeError('No output data. Scanned PDFs need OCR, which this version does not include.')
        return _publish(files, dest)


def merge_pdfs(paths, destination=None, cancel=None):
    import fitz
    paths = [Path(p).resolve() for p in paths]
    if len(paths) < 2 or any(p.suffix.lower() != '.pdf' for p in paths):
        raise ValueError('Choose at least two PDF files. Queue order is merge order.')
    dest = Path(destination).resolve() if destination else paths[0].parent / 'Velora Output'
    dest.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.velora-', dir=dest) as tmp:
        target = Path(tmp) / 'Merged.pdf'
        with fitz.open() as result:
            for p in paths:
                if cancel and cancel.is_set():
                    raise RuntimeError('Cancelled')
                with fitz.open(p) as doc:
                    if doc.needs_pass:
                        raise ValueError('Unlock protected PDFs before merging.')
                    result.insert_pdf(doc)
            result.save(target, deflate=True)
        if cancel and cancel.is_set():
            raise RuntimeError('Cancelled')
        return _publish([target], dest)
