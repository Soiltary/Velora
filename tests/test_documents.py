import hashlib
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
import fitz
from docx import Document
from PIL import Image
from velora import engine
from velora.documents import docx_to_pdf
from velora.i18n import set_language, tr

class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)
    def tearDown(self):
        set_language('en')
        self.tmp.cleanup()
    def test_pdf_to_editable_docx_and_original_preserved(self):
        src=self.root/'input.pdf'
        with fitz.open() as doc:
            doc.new_page().insert_text((50,50),'Editable PDF text')
            doc.new_page().insert_text((50,50),'Second PDF page')
            doc.save(src)
        digest=hashlib.sha256(src.read_bytes()).digest()
        out=engine.convert(src,'docx')[0]
        word=Document(out)
        text='\n'.join(p.text for p in word.paragraphs)
        self.assertIn('Editable PDF text',text)
        self.assertIn('Second PDF page',text)
        self.assertTrue(word.element.xpath('.//w:br[@w:type="page"]'))
        self.assertEqual(digest,hashlib.sha256(src.read_bytes()).digest())
    def test_scan_retained_as_image(self):
        image=self.root/'scan.png'
        Image.new('RGB',(100,100),'purple').save(image)
        src=self.root/'scan.pdf'
        with fitz.open() as doc:
            page=doc.new_page();page.insert_image(page.rect,filename=str(image));doc.save(src)
        out=engine.convert(src,'docx')[0]
        self.assertEqual(len(Document(out).inline_shapes),1)
    def test_docx_without_office_preserves_text_table_image(self):
        src=self.root/'input.docx';doc=Document();doc.add_paragraph('Native PDF export')
        doc.add_paragraph('مرحبا سلطان')
        table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Table cell';table.cell(0,1).text='Value'
        image=self.root/'inline.png';Image.new('RGB',(100,50),'purple').save(image);doc.add_picture(str(image));doc.save(src)
        with patch('velora.engine.libreoffice',return_value=None):
            out=engine.convert(src,'pdf')[0]
        with fitz.open(out) as pdf:
            self.assertEqual(len(pdf),1)
            self.assertFalse(pdf[0].get_text().strip())
            self.assertTrue(pdf[0].get_images())
            pix=pdf[0].get_pixmap();self.assertGreater(len(set(pix.samples)),5)
    def test_cancelled_conversion_publishes_nothing(self):
        src=self.root/'input.docx';doc=Document();doc.add_paragraph('Hello');doc.save(src)
        cancel=threading.Event();cancel.set()
        with patch('velora.engine.libreoffice',return_value=None):
            with self.assertRaisesRegex(RuntimeError,'Cancelled'):engine.convert(src,'pdf',cancel=cancel)
        self.assertFalse(list((self.root/'Velora Output').glob('*.pdf')))
    def test_protected_pdf_rejected(self):
        src=self.root/'locked.pdf'
        with fitz.open() as pdf:
            pdf.new_page();pdf.save(src,encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='password')
        with self.assertRaisesRegex(ValueError,'Password'):engine.convert(src,'docx')
    def test_language_switch_round_trip(self):
        set_language('ar');self.assertEqual(tr('Support'),'دعم')
        english='Finished: 2 succeeded, 1 failed or cancelled.'
        translated=tr(english);self.assertIn('نجح',translated)
        set_language('en');self.assertEqual(tr(translated),english)

if __name__=='__main__':unittest.main()
