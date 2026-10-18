import csv
import hashlib
import sys
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path

import fitz
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from velora import engine


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.image = self.root/'صورة test.png'
        Image.new('RGBA',(96,64),(120,60,240,128)).save(self.image)

    def tearDown(self):
        self.tmp.cleanup()

    def test_images_original_and_collision(self):
        before = hashlib.sha256(self.image.read_bytes()).digest()
        for target in ['png','jpg','webp','avif','bmp','tiff','ico','pdf','resize','rotate','grayscale','strip-metadata','compress']:
            result = engine.convert(self.image,target,options={'width':32,'height':32})
            self.assertTrue(result[0].stat().st_size)
            if target != 'pdf':
                with Image.open(result[0]) as im:
                    im.verify()
        a = engine.convert(self.image,'jpg')[0]
        b = engine.convert(self.image,'jpg')[0]
        self.assertNotEqual(a,b)
        self.assertEqual(before,hashlib.sha256(self.image.read_bytes()).digest())

    def test_pdf_pages_text_merge_split(self):
        pdf = self.root/'two.pdf'
        with fitz.open() as doc:
            for t in ['First page','Second page']:
                doc.new_page().insert_text((50,50),t)
            doc.save(pdf)
        self.assertEqual(len(engine.convert(pdf,'png')),2)
        self.assertEqual(len(engine.convert(pdf,'jpg')),2)
        split = engine.convert(pdf,'split-pages')
        merged = engine.merge_pdfs(split)
        with fitz.open(merged[0]) as doc:
            self.assertEqual(len(doc),2)
        self.assertIn('Second page',engine.convert(pdf,'txt')[0].read_text(encoding='utf-8'))
        self.assertTrue(engine.convert(pdf,'compress')[0].exists())

    def test_subtitle_timing_roundtrip(self):
        srt = self.root/'sub.srt'
        srt.write_text('1\n00:01:02,345 --> 00:01:05,678\nمرحبا بالعالم\n\n',encoding='utf-8')
        vtt = engine.convert(srt,'vtt')[0]
        back = engine.convert(vtt,'srt')[0]
        self.assertIn('00:01:02,345 --> 00:01:05,678',back.read_text(encoding='utf-8'))
        self.assertIn('مرحبا بالعالم',back.read_text(encoding='utf-8'))

    def test_csv_xlsx_no_formula_and_sheets(self):
        from openpyxl import load_workbook
        path = self.root/'sample.csv'
        path.write_text('Name,Value\nسلطان,=1+2\n',encoding='utf-8')
        xlsx = engine.convert(path,'xlsx')[0]
        wb = load_workbook(xlsx)
        self.assertEqual(wb.active['B2'].data_type,'s')
        wb.create_sheet('Other').append(['Second sheet'])
        wb.save(xlsx)
        wb.close()
        tsv = engine.convert(xlsx,'tsv')
        self.assertEqual(len(tsv),2)
        self.assertIn('=1+2',tsv[0].read_text(encoding='utf-8-sig'))

    def test_text_docx_html(self):
        path = self.root/'text.txt'
        path.write_text('مرحبا\nHello <world>',encoding='utf-8')
        doc = engine.convert(path,'docx')[0]
        txt = engine.convert(doc,'txt')[0]
        self.assertIn('مرحبا',txt.read_text(encoding='utf-8'))
        h = engine.convert(path,'html')[0]
        self.assertIn('&lt;world&gt;',h.read_text(encoding='utf-8'))

    def test_folder_zip_and_universal_zip(self):
        folder = self.root/'Archive'
        folder.mkdir()
        (folder/'a.txt').write_text('hello')
        archive = engine.convert(folder,'zip')[0]
        with zipfile.ZipFile(archive) as z:
            self.assertEqual(z.read('Archive/a.txt'),b'hello')
        self.assertTrue(engine.convert(self.image,'zip')[0].exists())
        with self.assertRaises(ValueError):
            engine.convert(folder,'zip',folder/'nested')

    def test_cancel_no_partial(self):
        event = threading.Event()
        event.set()
        with self.assertRaises(RuntimeError):
            engine.convert(self.image,'jpg',cancel=event)
        self.assertEqual(list((self.root/'Velora Output').iterdir()),[])

    def test_media_actual_transcode(self):
        video = self.root/'clip.mp4'
        engine.run([engine.ffmpeg(),'-nostdin','-y','-f','lavfi','-i','color=c=purple:s=64x64:r=12',
                    '-f','lavfi','-i','sine=frequency=440:sample_rate=44100','-t','1','-c:v','libx264','-c:a','aac',video])
        for target in ['mp4','mkv','webm','mov','gif','mp3','wav','compress','trim','extract-audio']:
            output = engine.convert(video,target,options={'start':0,'end':.5})
            self.assertGreater(output[0].stat().st_size,100)
        wav = engine.convert(video,'wav')[0]
        for target in ['flac','m4a','ogg','opus','aiff','mp3']:
            self.assertGreater(engine.convert(wav,target)[0].stat().st_size,100)


if __name__ == '__main__':
    unittest.main()
